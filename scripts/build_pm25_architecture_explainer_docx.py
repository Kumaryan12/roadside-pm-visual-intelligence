from __future__ import annotations

import os
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "deliverables"
OUT_PATH = OUT_DIR / "scientific_explanation_pm25_candidate_architecture.docx"
FIG_PATH = OUT_DIR / "pm25_candidate_architecture_clean.png"

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "EAF2F8"
ORANGE = "C65911"
LIGHT_ORANGE = "FCE4D6"
GREEN = "3B7D23"
LIGHT_GREEN = "E2F0D9"
PURPLE = "7030A0"
LIGHT_PURPLE = "EDE4F5"
INK = "1F2937"
MUTED = "5B6573"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E0E7"
RED = "9B1C1C"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(
    run,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr, fld_char_2])


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    before: float = 0,
    after: float = 6,
    line_spacing: float = 1.10,
    align=None,
    keep_with_next: bool = False,
):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep_with_next
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, size=11, color=INK)
    return p


def add_rich_paragraph(doc: Document, segments, *, before=0, after=6, line_spacing=1.10, align=None):
    p = add_paragraph(doc, before=before, after=after, line_spacing=line_spacing, align=align)
    for text, kwargs in segments:
        run = p.add_run(text)
        set_run_font(run, size=kwargs.pop("size", 11), color=kwargs.pop("color", INK), **kwargs)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    # python-docx may leave an empty run generated by the style; the explicit run
    # ensures consistent rendering across Word and LibreOffice.
    set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else NAVY, bold=True)
    return p


def add_equation(doc: Document, text: str, *, color: str = NAVY):
    p = add_paragraph(doc, before=4, after=8, line_spacing=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(text)
    set_run_font(run, name="Cambria Math", size=13, color=color, bold=True)
    return p


def add_callout(
    doc: Document,
    label: str,
    body: str,
    *,
    fill: str,
    accent: str,
    trailing_space: bool = True,
):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent_dxa=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=11, color=accent, bold=True)
    r = p.add_run(body)
    set_run_font(r, size=11, color=INK)
    if trailing_space:
        add_paragraph(doc, after=2)
    return table


def add_component_table(doc: Document) -> None:
    rows = [
        (
            "External CAMS background",
            "Regional atmospheric PM₂.₅ baseline aligned by time and route location.",
            "Represents transported and broad-scale pollution; it cannot resolve individual roadside events.",
        ),
        (
            "ResNet50 + GRU",
            "A seven-frame lens-6 sequence encoded to 2,048-D frame embeddings and processed in temporal order.",
            "Models short-term visual context and how the roadside scene evolves over roughly one minute.",
        ),
        (
            "YOLO traffic",
            "Vehicle classes, counts, occupancy, size/proximity and activity proxies.",
            "Makes traffic composition explicit; these are predictors, not measured vehicle PM contributions.",
        ),
        (
            "Road appearance",
            "Segmented-road area, brightness, saturation, texture, edge, dry/gray and related appearance indicators.",
            "Acts as a road-condition and resuspension-potential proxy, not a measurement of road-dust mass.",
        ),
        (
            "Meteorology / sensor / mobility",
            "Temperature, RH, gases, time, location, speed, course and related non-PM context.",
            "Represents dilution, dispersion, aerosol/sensor response and changing exposure geometry.",
        ),
        (
            "AlphaEarth",
            "A 64-dimensional annual learned geospatial embedding.",
            "Provides persistent wider-area context; individual dimensions do not have direct physical labels.",
        ),
        (
            "OSM (tested ablation)",
            "Road hierarchy and nearby activity/land-use counts within 250 m.",
            "Supplies off-camera context. It did not improve the selected final ensemble, so it is not in the frozen no-OSM candidate.",
        ),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [2200, 3260, 3900], indent_dxa=120)
    headers = ("Information source", "What enters the model", "Scientific role")
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=10, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_idx, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        if row_idx % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, "F8FAFC")
        for col_idx, text in enumerate(values):
            p = cells[col_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(text)
            set_run_font(r, size=9.2, color=INK, bold=(col_idx == 0))
    add_paragraph(doc, after=2)


def add_claims_table(doc: Document) -> None:
    rows = [
        ("Supported", "The model predicts roadside PM₂.₅ using an external regional baseline plus a learned local/sub-grid correction."),
        ("Supported", "YOLO, road, meteorology and geospatial variables provide candidate predictors of unresolved roadside variation."),
        ("Not supported", "The model separately measures vehicle-exhaust PM₂.₅, road-dust PM₂.₅ or construction PM₂.₅."),
        ("Not supported", "Feature importance proves causality or chemical source contribution."),
        ("Required next evidence", "Untouched future dates, raw one-second validation around peaks and independent source-specific measurements."),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2100, 7260], indent_dxa=120)
    for idx, text in enumerate(("Claim status", "Scientifically defensible interpretation")):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        r = cell.paragraphs[0].add_run(text)
        set_run_font(r, size=10, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    for status, text in rows:
        cells = table.add_row().cells
        fill = LIGHT_GREEN if status == "Supported" else (LIGHT_ORANGE if status == "Not supported" else LIGHT_PURPLE)
        set_cell_shading(cells[0], fill)
        r = cells[0].paragraphs[0].add_run(status)
        set_run_font(r, size=9.5, color=GREEN if status == "Supported" else (RED if status == "Not supported" else PURPLE), bold=True)
        r = cells[1].paragraphs[0].add_run(text)
        set_run_font(r, size=9.5, color=INK)
        for cell in cells:
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.05


def build_architecture_figure(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(13, 4.7), dpi=180)
    ax.set_xlim(0, 13)
    ax.set_ylim(0, 4.7)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    boxes = [
        (0.25, 0.55, 2.35, 3.55, "#EAF2F8", "#2E74B5", "1  EXTERNAL BACKGROUND",
         "Timestamp + route location\n↓\nCAMS regional PM₂.₅\n\nNot fitted to roadside\ntest targets"),
        (2.95, 0.55, 3.55, 3.55, "#EAF2F8", "#2E74B5", "2  TEMPORAL BRANCH",
         "7 lens-6 frames\n↓\nFrozen ResNet50 embeddings\n↓\nGRU local-residual estimate\n+\nRF residual correction"),
        (6.85, 0.55, 2.75, 3.55, "#FCE4D6", "#C65911", "3  TABULAR BRANCH",
         "YOLO traffic\n+ road appearance\n+ AlphaEarth\n(OSM tested separately)\n↓\nExtraTrees local-residual\nestimate"),
        (9.95, 0.55, 2.75, 3.55, "#EDE4F5", "#7030A0", "4  FUSION",
         "Validation-selected λ\n↓\nλ·temporal + (1−λ)·tabular\n↓\nAdd CAMS background\n↓\nFinal roadside PM₂.₅"),
    ]
    for x, y, w, h, fill, edge, title, body in boxes:
        patch = FancyBboxPatch(
            (x, y), w, h, boxstyle="round,pad=0.025,rounding_size=0.08",
            facecolor=fill, edgecolor=edge, linewidth=1.6
        )
        ax.add_patch(patch)
        ax.text(x + 0.16, y + h - 0.35, title, fontsize=10.5, color=edge, weight="bold", va="top")
        ax.text(x + w / 2, y + h / 2 - 0.15, body, fontsize=9.4, color="#1F2937",
                ha="center", va="center", linespacing=1.35)
    for x1, x2 in ((2.60, 2.95), (6.50, 6.85), (9.60, 9.95)):
        ax.add_patch(FancyArrowPatch((x1, 2.33), (x2, 2.33), arrowstyle="-|>", mutation_scale=14,
                                     linewidth=1.5, color="#667085"))
    ax.text(
        6.5, 0.20,
        "Training target for both local branches: Δlocal = roadside sensor PM₂.₅ − external CAMS background",
        ha="center", va="center", fontsize=10.3, color="#17365D", weight="bold"
    )
    plt.tight_layout(pad=0.5)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.40)
    section.footer_distance = Inches(0.40)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, NAVY, 8, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("TECHNICAL NOTE  |  ROADSIDE PM₂.₅ PREDICTION")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_after = Pt(0)
    add_page_number(fp)


def build_document() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_architecture_figure(FIG_PATH)

    doc = Document()
    configure_styles(doc)

    # Page 1: memo masthead / technical note
    p = add_paragraph(doc, "SCIENTIFIC ARCHITECTURE NOTE", before=20, after=4)
    set_run_font(p.runs[0], size=10, color=ORANGE, bold=True)

    p = add_paragraph(doc, after=5, line_spacing=1.0)
    r = p.add_run("Scientific Rationale and Operation of the\nPM₂.₅ Prediction Architecture")
    set_run_font(r, size=25, color=NAVY, bold=True)

    p = add_paragraph(doc, after=16)
    r = p.add_run("External atmospheric background with a nested temporal–tabular local-residual ensemble")
    set_run_font(r, size=13, color=MUTED, italic=True)

    meta = [
        ("Purpose", "Technical explanation for academic review"),
        ("Dataset context", "MUMMA five-day moving-platform roadside observations"),
        ("Target", "Roadside sensor PM₂.₅ concentration"),
        ("Validation principle", "Complete-date outer holdout; model choices use training/validation data only"),
        ("Prepared", "24 July 2026"),
    ]
    for label, value in meta:
        p = add_paragraph(doc, after=2, line_spacing=1.0)
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=10.5, color=NAVY, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=INK)

    add_paragraph(doc, after=3)
    add_callout(
        doc,
        "Central scientific idea",
        "A roadside monitor observes broad regional pollution plus unresolved local and sub-grid variation. "
        "The architecture uses an independent atmospheric product for the regional baseline and learns only "
        "the remaining roadside residual from visual, traffic, road, meteorological and geospatial information.",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    add_heading(doc, "Executive summary", 1)
    add_paragraph(
        doc,
        "The architecture is based on separation of spatial and temporal scales. CAMS supplies a coarse external "
        "estimate of regional atmospheric PM₂.₅. The roadside sensor may differ from that estimate because of "
        "fine-scale traffic activity, road-surface resuspension, street geometry, meteorological dispersion, "
        "unresolved sources, product bias and sensor error. Two complementary models predict this difference: "
        "a ResNet50–GRU temporal image branch with structured residual correction, and an ExtraTrees tabular/"
        "geospatial branch. Their predictions are combined using a weight selected on validation data, after "
        "which the external background is added back to obtain the final roadside PM₂.₅ estimate."
    )
    add_paragraph(
        doc,
        "The method is a predictive concentration model, not chemical source apportionment. Vehicle counts and "
        "road appearance are covariates associated with local conditions; they are not direct measurements of "
        "vehicle-exhaust or road-dust mass."
    )

    doc.add_page_break()

    # Page 2: first principles + figure
    add_heading(doc, "1. First-principles formulation", 1)
    add_paragraph(
        doc,
        "At a given location s and time t, the concentration measured by the moving roadside sensor can be "
        "represented conceptually as a regional atmospheric component, an unresolved local/sub-grid deviation, "
        "and remaining error:"
    )
    add_equation(doc, "y(s,t) = B(s,t) + Δlocal(s,t) + ε(s,t)")
    add_rich_paragraph(
        doc,
        [
            ("y(s,t)", {"bold": True, "color": NAVY}),
            (" is roadside PM₂.₅; ", {}),
            ("B(s,t)", {"bold": True, "color": BLUE}),
            (" is the external regional background; ", {}),
            ("Δlocal(s,t)", {"bold": True, "color": ORANGE}),
            (" is the residual at scales unresolved by the background product; and ", {}),
            ("ε(s,t)", {"bold": True, "color": PURPLE}),
            (" contains measurement error and unobserved processes.", {}),
        ],
        after=8,
    )
    add_paragraph(
        doc,
        "The training target for both local branches is therefore the measured roadside value minus the aligned "
        "CAMS estimate:"
    )
    add_equation(doc, "Δlocal(s,t) = yroadside(s,t) − yCAMS(s,t)", color=ORANGE)
    add_callout(
        doc,
        "Important interpretation",
        "Δlocal is not a pure local-emissions concentration. It may contain genuine near-road enhancement, "
        "negative local deviations, CAMS bias, coarse-grid mismatch, timing error, sensor response and other "
        "unobserved effects. “Local/sub-grid residual” is the scientifically safest term.",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )

    p = add_paragraph(doc, "Figure 1. Operational flow of the candidate architecture.", before=6, after=4)
    set_run_font(p.runs[0], size=9.5, color=MUTED, italic=True)
    doc.add_picture(str(FIG_PATH), width=Inches(6.5))
    picture_properties = doc.inline_shapes[-1]._inline.docPr
    picture_properties.set(
        "descr",
        "Flow diagram showing an external CAMS background, a ResNet50-GRU temporal local-residual branch "
        "with Random-Forest correction, an ExtraTrees tabular local-residual branch, and validation-selected "
        "fusion followed by reconstruction of roadside PM2.5.",
    )
    picture_properties.set("title", "Candidate roadside PM2.5 prediction architecture")
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(6)

    doc.add_page_break()

    # Page 3: external + temporal
    add_heading(doc, "2. External atmospheric background", 1)
    add_paragraph(
        doc,
        "CAMS is an atmospheric composition modelling system. It combines emissions inventories, atmospheric "
        "transport, meteorology, chemistry, deposition, aerosol processes and assimilated observations. It is "
        "not a direct satellite photograph of ground-level PM₂.₅. For each roadside timestamp, the corresponding "
        "coarse regional estimate is aligned and used as B(s,t). CAMS is not fitted to the held-out MUMMA targets."
    )
    add_paragraph(
        doc,
        "Because its spatial grid and temporal update interval are much coarser than the roadside measurements, "
        "CAMS cannot resolve the vehicle immediately ahead, a single junction, a local construction site, "
        "street-canyon circulation or second-to-second fluctuations. These are exactly the scales delegated to "
        "the local branches."
    )
    add_callout(
        doc,
        "Why this is not target leakage",
        "The background is generated independently of the roadside test target and is available at inference "
        "time. It would become leakage only if test-day roadside PM₂.₅ were used to calibrate or modify the "
        "background before prediction.",
        fill=LIGHT_GREEN,
        accent=GREEN,
    )

    add_heading(doc, "3. Temporal local-residual branch", 1)
    add_heading(doc, "3.1 Seven-frame visual sequence", 2)
    add_paragraph(
        doc,
        "The temporal branch receives seven consecutive lens-6 images. With approximately ten-second sampling, "
        "the sequence describes roughly one minute of recent roadside context. Sequences are kept inside a "
        "single run, preventing a sequence from crossing into another run or day."
    )
    add_heading(doc, "3.2 Frozen ResNet50 representation", 2)
    add_equation(doc, "zt = ResNet50(It),   zt ∈ ℝ²⁰⁴⁸")
    add_paragraph(
        doc,
        "A pretrained, frozen ResNet50 converts each image into a 2,048-dimensional embedding. The embedding "
        "summarizes general scene appearance, including traffic, road surface, built environment, visibility and "
        "illumination. It does not directly measure PM₂.₅. Freezing the backbone reduces the number of parameters "
        "that must be estimated from only five days of PM-labelled observations."
    )
    add_heading(doc, "3.3 GRU temporal encoder", 2)
    add_equation(doc, "Δ̂temp = fGRU(zt−6, …, zt)")
    add_paragraph(
        doc,
        "A one-layer GRU with hidden size 128 processes the embeddings in chronological order. It can distinguish "
        "a persistent nearby heavy vehicle from a single-frame detection, sustained congestion from a transient "
        "event, and gradual scene change from abrupt change. Its output is the image-based estimate of the local "
        "residual."
    )

    doc.add_page_break()

    # Page 4: structured features + residual correction
    add_heading(doc, "4. Structured information and why it is still useful", 1)
    add_paragraph(
        doc,
        "The camera contains some of the same information implicitly, but the dataset is too small to assume that "
        "a generic image embedding will reliably discover every PM-relevant quantity. Structured variables act "
        "as inductive biases: they make selected scene properties explicit and add information beyond the camera’s "
        "field of view."
    )
    add_component_table(doc)

    add_heading(doc, "5. Random-Forest correction of temporal errors", 1)
    add_paragraph(
        doc,
        "The GRU can make systematic errors. Its residual on an eligible training or validation observation is:"
    )
    add_equation(doc, "rt = Δlocal,t − Δ̂temp,t")
    add_paragraph(
        doc,
        "A Random Forest predicts this remaining error from non-PM sensor, meteorological, mobility, YOLO and road "
        "features. The corrected temporal estimate is:"
    )
    add_equation(doc, "Δ̂corr,temp = Δ̂temp + r̂RF")
    add_paragraph(
        doc,
        "This structure asks a narrower question than direct prediction: given what the image sequence already "
        "predicted, do the structured conditions indicate a systematic adjustment? Correction predictions are "
        "cross-fitted by run so that validation residuals are not produced by a correction model trained on the "
        "same run."
    )

    doc.add_page_break()

    # Page 5: tabular + fusion
    add_heading(doc, "6. Independent tabular/geospatial branch", 1)
    add_paragraph(
        doc,
        "The second branch predicts the same local residual independently using explicit scene and spatial "
        "information. In the selected no-OSM candidate, the principal inputs are YOLO traffic, road appearance "
        "and the 64-dimensional AlphaEarth embedding:"
    )
    add_equation(doc, "Δ̂tab = fExtraTrees(XYOLO, Xroad, XAlphaEarth)")
    add_paragraph(
        doc,
        "ExtraTrees is a randomized ensemble of decision trees. It can represent nonlinear thresholds and "
        "interactions without assuming a linear relationship. For example, heavy-vehicle activity may have a "
        "different statistical association under different road appearance and neighbourhood conditions."
    )
    add_callout(
        doc,
        "Diagram terminology correction",
        "ExtraTrees is not gradient boosting. The architecture diagram should label this block “ExtraTrees "
        "ensemble regressor.” Gradient boosting was evaluated as a separate candidate model.",
        fill=LIGHT_ORANGE,
        accent=ORANGE,
    )

    add_heading(doc, "7. Validation-selected fusion", 1)
    add_paragraph(
        doc,
        "The temporal and tabular branches are complementary: the first represents recent visual dynamics, while "
        "the second emphasizes explicit traffic, road and persistent spatial context. Their local-residual "
        "predictions are combined using a convex weight:"
    )
    add_equation(doc, "Δ̂fused = λΔ̂corr,temp + (1 − λ)Δ̂tab,   0 ≤ λ ≤ 1", color=PURPLE)
    add_paragraph(
        doc,
        "The value of λ is selected using validation RMSE only and then frozen before the outer test date is "
        "evaluated. λ = 1 uses only the temporal branch; λ = 0 uses only the tabular branch. A fixed 50:50 average "
        "is a useful diagnostic, but it is not the primary result when the prescribed selection rule chooses λ "
        "from validation data."
    )
    add_heading(doc, "8. Final roadside PM₂.₅ reconstruction", 1)
    add_equation(
        doc,
        "ŷPM₂.₅ = yCAMS + λΔ̂corr,temp + (1 − λ)Δ̂tab",
        color=GREEN,
    )
    add_paragraph(
        doc,
        "For example, if CAMS gives 70 µg/m³, the corrected temporal branch predicts a residual of 35 µg/m³, "
        "the tabular branch predicts 25 µg/m³ and λ = 0.7, the fused residual is 32 µg/m³ and the final prediction "
        "is 102 µg/m³."
    )

    doc.add_page_break()

    # Page 6: objections / scientific reasoning
    add_heading(doc, "9. Responses to the key scientific questions", 1)
    add_heading(doc, "9.1 If road dust is already in the atmosphere, why include road features?", 2)
    add_paragraph(
        doc,
        "CAMS may contain regional suspended dust at coarse scale, but it cannot resolve a new near-road flux "
        "created when local traffic resuspends material from a particular road surface. Subtracting CAMS removes "
        "the broad component it represents; road appearance and traffic variables are then offered as predictors "
        "of the unresolved difference. Nevertheless, the road variables are proxies only and cannot identify "
        "road-dust mass without chemical or particle-size reference measurements."
    )
    add_heading(doc, "9.2 If the image already sees vehicles and the road, why add YOLO and road variables?", 2)
    add_paragraph(
        doc,
        "A generic ResNet embedding encodes the scene implicitly, but with limited labelled data it may not "
        "reliably learn that particular regions correspond to trucks, vehicle occupancy or road texture. YOLO and "
        "road segmentation expose those quantities directly. Their incremental value must still be demonstrated "
        "through held-out-date ablation; redundancy is possible and more features do not automatically improve "
        "generalization."
    )
    add_heading(doc, "9.3 Why include OSM or AlphaEarth?", 2)
    add_paragraph(
        doc,
        "A camera has a narrow instantaneous field of view. OSM and AlphaEarth can represent broader, persistent "
        "context such as road hierarchy, commercial activity and urban form that may be behind the vehicle or "
        "outside the image. In the actual experiments, OSM did not improve the selected nested ensemble, so it "
        "should be reported as an ablation rather than a required component."
    )
    add_heading(doc, "9.4 Does the model estimate vehicle PM₂.₅ and road-dust PM₂.₅ separately?", 2)
    add_paragraph(
        doc,
        "No. The only supervised label is total roadside sensor PM₂.₅. Therefore, separate vehicle, road-dust and "
        "background source fractions are not identifiable. Such a claim would require source-specific evidence "
        "such as chemical tracers, differential OPC bins, black carbon, independent background monitors or "
        "validated source profiles."
    )
    add_heading(doc, "9.5 Are feature importances causal?", 2)
    add_paragraph(
        doc,
        "No. Traffic, time, route, weather, land use and daily background are correlated. Feature importance or "
        "correlation indicates predictive association under the observed data distribution; it does not prove "
        "that changing the feature would cause the corresponding PM₂.₅ change."
    )

    doc.add_page_break()

    # Page 7: validation + claims + concise summary
    add_heading(doc, "10. Fair evaluation procedure", 1)
    add_paragraph(
        doc,
        "Each of the five dates is held out in turn as an outer test date. The remaining dates are used for model "
        "training and validation. Image sequences remain within run boundaries; preprocessing and model fitting "
        "use outer-training data only; correction is cross-fitted; and λ is chosen using validation predictions. "
        "Only after these choices are frozen is the complete held-out date evaluated."
    )
    add_callout(
        doc,
        "Current evidence level",
        "Whole-date testing is more credible than random overlapping-window testing, but the architecture was "
        "developed after repeated inspection of the same five dates. It remains an exploratory candidate until "
        "it is frozen and tested once on untouched future dates.",
        fill=LIGHT_PURPLE,
        accent=PURPLE,
    )

    add_heading(doc, "11. Scientifically defensible interpretation", 1)
    add_claims_table(doc)

    add_heading(doc, "12. Concise explanation for discussion", 1)
    add_paragraph(
        doc,
        "The architecture models roadside PM₂.₅ by separating regional and local scales. CAMS supplies an "
        "independently generated coarse atmospheric background. The model then predicts the difference between "
        "the roadside sensor and that background. A ResNet50–GRU branch estimates this local/sub-grid residual "
        "from a recent image sequence, while structured traffic, road, meteorological and mobility variables "
        "correct systematic temporal-model errors. An independent ExtraTrees branch estimates the same residual "
        "from explicit visual and geospatial context. The two estimates are combined with a validation-selected "
        "weight, and the CAMS background is added back to obtain final roadside PM₂.₅. Vehicle and road variables "
        "are predictive proxies rather than measured source concentrations, so the architecture predicts total "
        "roadside concentration but does not claim chemical source apportionment."
    )

    # Keep document metadata neutral and ready to share.
    doc.core_properties.title = "Scientific Rationale and Operation of the PM2.5 Prediction Architecture"
    doc.core_properties.subject = "External atmospheric background and nested temporal-tabular local-residual ensemble"
    doc.core_properties.author = "Roadside PM Visual Intelligence Project"
    doc.core_properties.keywords = "PM2.5, CAMS, ResNet50, GRU, ExtraTrees, roadside air pollution"
    doc.core_properties.comments = "Prepared for academic technical review."

    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    result = build_document()
    print(result)
