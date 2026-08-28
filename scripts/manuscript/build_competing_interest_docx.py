"""Build the separate Elsevier competing-interest declaration."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = (
    ROOT
    / "docs"
    / "environmental_modelling_software_submission_v1"
    / "declaration_of_competing_interest.docx"
)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(20)
    run = title.add_run("Declaration of Competing Interest")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(18)

    table = document.add_table(rows=2, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(4.95)
    labels = [
        (
            "Manuscript title",
            "A Leakage-Aware Software Framework for Background-Conditioned PM2.5 "
            "Prediction under Grouped Temporal and Spatial Shifts",
        ),
        ("Author", "Aryan Satyendra Kumar"),
    ]
    for row, (label, value) in zip(table.rows, labels):
        row.cells[0].width = Inches(1.55)
        row.cells[1].width = Inches(4.95)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.bold = True
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.add_run(value)

    document.add_paragraph()
    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(10)
    heading_run = heading.add_run("Declaration")
    heading_run.bold = True
    heading_run.font.size = Pt(13)

    statement = document.add_paragraph()
    statement.paragraph_format.space_after = Pt(24)
    statement.add_run(
        "The author declares that there are no known competing financial interests "
        "or personal relationships that could have appeared to influence the work "
        "reported in this paper."
    )

    document.add_paragraph("Name: Aryan Satyendra Kumar")
    document.add_paragraph("Signature: ____________________________________")
    document.add_paragraph("Date: 29 August 2026")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Environmental Modelling & Software submission")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)

    document.core_properties.title = "Declaration of Competing Interest"
    document.core_properties.author = "Aryan Satyendra Kumar"
    document.core_properties.subject = "Environmental Modelling & Software submission"
    document.core_properties.comments = "Generated from the repository submission source."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
